"""File BIR zonal value workbooks from raw/bir/_inbox into raw/bir/RDO-<code>/ and write raw/bir/inventory.csv.

Everything is read from the workbook contents, never the file or folder name:
  - the RDO code and name come from the NOTICE sheet ("RDO No. 47 - East Makati City is composed of ..."),
    cross-checked against the subject line of the newest Department Order sheet;
  - the Department Orders come from the NOTICE table (revision, effectivity dates, LGUs covered) and from the
    header of each DO sheet ("DEPARTMENT ORDER NO. 037 - 2021"). The NOTICE "See details on" column is not
    trusted: it points at the wrong sheet in several files.

Usage:
  python -m pipeline.inventory            # file new inbox files, then rebuild the inventory from all RDO-* folders
Outputs:
  raw/bir/inventory.csv                   one row per (file, Department Order)
  raw/bir/filing_log.txt                  every move, duplicate and unreadable file, appended run after run
  raw/bir/inventory_log.txt               notes on the current files: anything that needed judgment
  pipeline/review/bir_sheets.csv          one row per DO sheet: which sheet holds which DO, revision per LGU
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import xlrd
from rapidfuzz import fuzz, process

from .normalize import ascii_fold, norm_city

ROOT = Path(__file__).resolve().parent.parent
BIR = ROOT / 'raw' / 'bir'
INBOX = BIR / '_inbox'
UNREADABLE = BIR / '_unreadable'
DUPLICATES = BIR / '_duplicates'
REVIEW = Path(__file__).resolve().parent / 'review'
CITIES = ROOT / 'public' / 'data' / 'cities.geojson'
IN_SCOPE_PROVINCES = {'Metro Manila', 'Cavite', 'Laguna', 'Rizal', 'Bulacan'}
EXCEL_SUFFIXES = {'.xls', '.xlsx', '.xlsm'}
HEADER_ROWS = 40  # DO number, date and subject sit in the first rows of each DO sheet

# BIR spellings that fuzzy matching should not have to guess
LGU_ALIASES = {
    'montalban': 'Rodriguez', 'rodriguezmontalban': 'Rodriguez', 'mendeznunez': 'Mendez',
    'damarinas': 'Dasmariñas', 'trecemartirez': 'Trece Martires', 'gma': 'Gen. Mariano Alvarez',
    'generalmarianoalvarez': 'Gen. Mariano Alvarez', 'genmarianoalvarez': 'Gen. Mariano Alvarez',
    'genemilioaguinaldo': 'General Emilio Aguinaldo', 'baliuag': 'Baliwag', 'obandoo': 'Obando',
    'stamaria': 'Santa Maria', 'starosa': 'Santa Rosa', 'stacruz': 'Santa Cruz', 'gentrias': 'General Trias',
    'generaleaguinaldo': 'General Emilio Aguinaldo', 'geneaguinaldo': 'General Emilio Aguinaldo', 'bulakan': 'Bulacan',
}


# ---------------------------------------------------------------- workbook access

class Book:
    """Minimal read-only view over .xls (xlrd) and .xlsx (openpyxl): sheet names and the first N rows."""

    def __init__(self, path: Path):
        self.path = path
        self.xlsx = path.suffix.lower() in ('.xlsx', '.xlsm')
        if self.xlsx:
            self.wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.names = self.wb.sheetnames
        else:
            self.wb = xlrd.open_workbook(path, on_demand=True)
            self.names = self.wb.sheet_names()

    def rows(self, name: str, limit: int) -> list[list]:
        if self.xlsx:
            ws = self.wb[name]
            return [list(r) for _, r in zip(range(limit), ws.iter_rows(values_only=True))]
        sh = self.wb.sheet_by_name(name)
        out = []
        for i in range(min(limit, sh.nrows)):
            row = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, self.wb.datemode))
                else:
                    row.append(c.value)
            out.append(row)
        return out

    def close(self):
        if self.xlsx:
            self.wb.close()
        else:
            self.wb.release_resources()


def text(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'\s+', ' ', str(v)).strip()


def line(row: list) -> str:
    return ' '.join(t for t in (text(v) for v in row) if t)


# ---------------------------------------------------------------- parsing helpers

# years are \d{4}|\d{2} in that order: "059-2023" must not read as year "20"
DO_RE = re.compile(r'^\s*(?:D[O0]\.?\s*(?:No\.?)?\s*)?(\d{1,3})\s*-\s*(\d{4}|\d{2})\s*$', re.I)
DO_HEADER_RE = re.compile(r'DEPARTMENT\s+ORDER\s+N[O0]\.?\s*(\d{1,3})\s*-\s*(\d{4}|\d{2})\b', re.I)
DO_IN_NAME_RE = re.compile(r'(\d{1,3})\s*-\s*(\d{4}|\d{2})\b')
RDO_RE = re.compile(r'\bRDO\s*(?:N[O0]\.?|#)?\s*(\d{1,2})\s*-?\s*([A-C])?\b\s*[-–]?\s*(.*)', re.I)
RDO_SUBJECT_RE = re.compile(r'REVENUE\s+DISTRICT\s+(?:OFFICE\s+)?(?:\(RDO\)\s*)?N[O0]\.?\s*(\d{1,2})\s*-?\s*([A-C])?\b', re.I)


def year4(y: str) -> int:
    n = int(y)
    if len(y) == 4:
        return n
    return 1900 + n if n >= 50 else 2000 + n


def canon_do(num: str, yr: str) -> str:
    return f'DO {int(num)}-{year4(yr)}'


def parse_do(s) -> str | None:
    m = DO_RE.match(text(s))
    return canon_do(*m.groups()) if m else None


def rdo_code(num: str, letter: str | None) -> str:
    letter = (letter or '').upper()
    return f'{int(num):02d}{letter}' if letter else f'{int(num):03d}'


MONTHS = {m: i for i, m in enumerate(['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'], 1)}


def parse_date(v) -> date | None:
    """Excel serials, datetimes, and BIR's assorted text dates ('4-June-21', '25-Sept-94', '15-10-2017')."""
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        if 20000 < v < 60000:  # Excel serial date
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        return None
    s = text(v)
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return date(*map(int, m.groups()))
    m = re.match(r'^(\d{1,2})[-\s/]([A-Za-z]{3,9})\.?[-\s/,]*(\d{2}|\d{4})$', s)
    if m and m.group(2)[:3].lower() in MONTHS:
        return date(year4(m.group(3)), MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
    m = re.match(r'^(\d{1,2})([-/])(\d{1,2})[-/](\d{4})$', s)
    if m:
        # slashes come from Excel's US export (9/17/2023 = Sept 17); dashes are written day first (15-10-2017)
        a, sep, b, y = int(m.group(1)), m.group(2), int(m.group(3)), int(m.group(4))
        mo, d = (a, b) if sep == '/' else (b, a)
        if mo > 12:
            mo, d = d, mo
        return date(y, mo, d)
    m = re.match(r'^([A-Za-z]{3,9})\.?\s+(\d{1,2}),?\s+(\d{4})$', s)
    if m and m.group(1)[:3].lower() in MONTHS:
        return date(int(m.group(3)), MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
    return None


def split_lgus(s: str) -> list[str]:
    s = re.sub(r'\b(Cities|City|Municipalities|Municipality) of\b', ',', s, flags=re.I)
    s = re.sub(r'\band\b|&|/', ',', s)
    out = []
    for part in s.split(','):
        p = part.strip(' .;:-')
        if p and not re.fullmatch(r'(the|following|of)', p, re.I):
            out.append(p)
    return out


# ---------------------------------------------------------------- gazetteer (in-scope LGUs)

class Gazetteer:
    def __init__(self):
        feats = json.loads(CITIES.read_text(encoding='utf-8'))['features']
        # several LGUs share a name across provinces (Santa Maria: Bulacan and Laguna), so keep them all
        self.entries: list[tuple[str, str]] = [(f['properties']['n'], f['properties']['p']) for f in feats]
        self.by_key: dict[str, tuple[str, str]] = {}
        for name, prov in self.entries:
            self.by_key.setdefault(self.key(name), (name, prov))
        self.ncr = {self.key(n): (n, p) for n, p in self.entries if p == 'Metro Manila'}

    @staticmethod
    def key(name: str) -> str:
        k = norm_city(name)
        k = re.sub(r'(city|municipality)$', '', k)
        return norm_city(LGU_ALIASES[k]) if k in LGU_ALIASES else k

    def lookup(self, raw: str) -> tuple[str, str] | None:
        k = self.key(raw)
        if k in self.by_key:
            return self.by_key[k]
        hit = process.extractOne(k, list(self.by_key), scorer=fuzz.ratio, score_cutoff=88)
        return self.by_key[hit[0]] if hit else None

    def ncr_cities_in(self, s: str) -> list[str]:
        t = ' ' + re.sub(r'[^a-z0-9]+', ' ', ascii_fold(s).lower()) + ' '
        found = []
        for name, _ in self.ncr.values():
            n = re.sub(r'[^a-z0-9]+', ' ', ascii_fold(name).lower()).strip()
            if f' {n} ' in t:
                found.append(name)
        return sorted(set(found))


# ---------------------------------------------------------------- per-file reading

@dataclass
class DOInfo:
    do: str
    revisions: list[str] = field(default_factory=list)
    froms: list[date] = field(default_factory=list)
    lgus_text: list[str] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    in_notice: bool = False


@dataclass
class FileInfo:
    path: Path
    rdo: str
    rdo_name: str
    region: str
    composition: str
    composed_of: str  # barangays | lgus | ''
    dos: dict[str, DOInfo]
    sheet_rows: list[dict]
    notes: list[str]


def find_notice(book: Book) -> str | None:
    for n in book.names:
        if n.strip().upper().replace(' ', '') == 'NOTICE':
            return n
    for n in book.names:  # fall back to any sheet whose first rows say NOTICE
        if any('NOTICE' in line(r).replace(' ', '').upper() for r in book.rows(n, 3)):
            return n
    return None


def read_notice(rows: list[list], notes: list[str]):
    """-> (rdo_code, rdo_name, region, composition text, composed_of, {do: DOInfo})."""
    lines = [line(r) for r in rows]
    region = next((l for l in lines if re.match(r'^RR\b', l, re.I)), '')
    rdo_line = next((l for l in lines if RDO_RE.search(l) and 'composed' in l.lower()), None)
    if rdo_line is None:
        rdo_line = next((l for l in lines if RDO_RE.search(l)
                         and not re.search(r'pursuant|splitted|formerly|merged', l, re.I)), None)
    rdo = name = None
    if rdo_line:
        m = RDO_RE.search(rdo_line)
        rdo = rdo_code(m.group(1), m.group(2))
        name = re.split(r'\s+is\s+composed', m.group(3), flags=re.I)[0].strip(' -:')

    # composition: lines after "composed of" up to the schedule table title
    comp, composed_of, grab = [], '', False
    for l in lines:
        if re.search(r'schedules? of zonal values', l, re.I):
            break
        if 'composed of' in l.lower():
            composed_of = 'barangays' if re.search(r'barangays\s*:?\s*$', l, re.I) else 'lgus'
            grab = True
            continue
        if grab and l:
            comp.append(l)
    composition = ' '.join(comp)

    # DO table
    hdr_i = next((i for i, r in enumerate(rows) if re.search(r'\brevision\b', line(r), re.I)), None)
    dos: dict[str, DOInfo] = {}
    if hdr_i is None:
        notes.append('NOTICE has no Department Order table')
        return rdo, name, region, composition, composed_of, dos
    hdr = [text(v).lower() for v in rows[hdr_i]]
    col = lambda pred: next((j for j, h in enumerate(hdr) if pred(h)), None)
    c_do = col(lambda h: h.startswith('do'))
    c_rev = col(lambda h: 'revision' in h)
    c_from = col(lambda h: 'effect' in h or 'effict' in h)
    c_cov = col(lambda h: any(w in h for w in ('covered', 'cit', 'municipal', 'barangay')))
    cov_is_brgy = c_cov is not None and 'barangay' in hdr[c_cov]
    cur = None
    for r in rows[hdr_i + 2:]:  # skip the From/To sub-header
        cells = r + [None] * (len(hdr) - len(r))
        if not line(r):
            continue
        do = parse_do(cells[c_do]) if c_do is not None else None
        if do:
            cur = dos.setdefault(do, DOInfo(do))
            cur.in_notice = True
        if cur is None:
            continue
        # a continuation row (no DO number) adds another LGU group, or wraps the previous LGU text
        rev = text(cells[c_rev]) if c_rev is not None else ''
        if rev:
            cur.revisions.append(rev)
        # any revision or date on the row means a new LGU group; otherwise the LGU text wrapped
        dated = c_from is not None and any(text(c) for c in cells[c_from:c_from + 2])
        d = parse_date(cells[c_from]) if c_from is not None else None
        if d:
            cur.froms.append(d)
        elif do and c_from is not None and text(cells[c_from]):
            notes.append(f'{do}: could not read effectivity date {text(cells[c_from])!r}')
        if c_cov is not None and not cov_is_brgy:
            cov = text(cells[c_cov])
            if cov and not parse_date(cells[c_cov]):
                if do or rev or dated or not cur.lgus_text:
                    cur.lgus_text.append(cov)
                else:
                    prev = cur.lgus_text[-1]
                    # "Mataas" + "na Kahoy" is one name; "Trece Martirez" + "Alfonso, ..." are two
                    sep = ' ' if prev.endswith(',') or cov[:1].islower() else ', '
                    cur.lgus_text[-1] = prev + sep + cov
    return rdo, name, region, composition, composed_of, dos


def read_file(path: Path, gaz: Gazetteer) -> FileInfo:
    notes: list[str] = []
    book = Book(path)
    try:
        notice = find_notice(book)
        if notice is None:
            raise ValueError('no NOTICE sheet')
        rdo, name, region, composition, composed_of, dos = read_notice(book.rows(notice, 200), notes)

        sheet_rows, subject_rdos = [], []
        for sname in book.names:
            if sname == notice:
                continue
            rows = book.rows(sname, HEADER_ROWS)
            body = [line(r) for r in rows]
            if not any(body):
                notes.append(f'sheet {sname!r} is empty; ignored')
                continue
            joined = ' '.join(body)
            m = DO_HEADER_RE.search(joined)
            do_hdr = canon_do(*m.groups()) if m else None
            m2 = DO_IN_NAME_RE.search(sname)
            do_name = canon_do(*m2.groups()) if m2 else None
            do = do_hdr or do_name
            if do is None:
                notes.append(f'sheet {sname!r}: no Department Order number found; ignored')
                continue
            if do_hdr and do_name and do_hdr != do_name:
                notes.append(f'sheet {sname!r}: tab says {do_name} but header says {do_hdr}; using header')
            subj = re.search(r'SUBJECT\s*:?(.*?)(?:\bTO\s*:|$)', joined, re.I | re.S)
            subject = subj.group(1) if subj else joined[:600]
            ms = RDO_SUBJECT_RE.search(subject) or RDO_SUBJECT_RE.search(joined)
            subj_rdo = rdo_code(ms.group(1), ms.group(2)) if ms else None
            info = dos.setdefault(do, DOInfo(do))
            info.sheets.append(sname)
            sheet_rows.append({'do': do, 'sheet': sname, 'do_from_header': do_hdr or '',
                               'do_from_tab': do_name or '', 'rdo_in_subject': subj_rdo or '',
                               'subject': re.sub(r'\s+', ' ', subject).strip()[:300]})
            subject_rdos.append((do, subj_rdo, subject))

        # RDO cross-check against the newest DO's subject line
        newest = max((d for d in dos.values() if d.sheets), key=lambda d: (int(d.do.split('-')[1]), int(d.do.split()[1].split('-')[0])), default=None)
        newest_rdo = next((r for d, r, _ in subject_rdos if newest and d == newest.do and r), None)
        if rdo is None and newest_rdo:
            rdo = newest_rdo
            notes.append(f'RDO taken from {newest.do} subject line (NOTICE had none)')
        elif rdo and newest_rdo and newest_rdo != rdo:
            notes.append(f'NOTICE says RDO {rdo} but {newest.do} subject says RDO {newest_rdo}')
        if rdo is None:
            raise ValueError('no RDO number in NOTICE or DO subject lines')
        if not name:
            name = ''
        for d in dos.values():
            if d.in_notice and not d.sheets:
                notes.append(f'{d.do}: listed in NOTICE but no sheet holds it')
            if d.sheets and not d.in_notice:
                notes.append(f'{d.do}: sheet {d.sheets} not listed in NOTICE; no effectivity date')
            if d.in_notice and not d.froms:
                notes.append(f'{d.do}: NOTICE gives no effectivity date')
            if len(set(d.froms)) > 1:
                notes.append(f'{d.do}: several effectivity dates by LGU {sorted(set(map(str, d.froms)))}; using the earliest')
        newest_subject = next((s for d, _, s in subject_rdos if newest and d == newest.do), '')
        fi = FileInfo(path, rdo, name, region, composition, composed_of, dos, sheet_rows, notes)
        fi.newest_subject = newest_subject  # type: ignore[attr-defined]
        return fi
    finally:
        book.close()


# ---------------------------------------------------------------- LGUs + scope

def resolve_lgus(raw_names: list[str], gaz: Gazetteer) -> tuple[list[str], Counter]:
    names, provs = [], Counter()
    for raw in raw_names:
        hit = gaz.lookup(raw)
        if hit:
            names.append(hit[0])
            provs[hit[1]] += 1
        else:
            names.append(raw)
            provs['(other)'] += 1
    return list(dict.fromkeys(names)), provs


def file_lgus(fi: FileInfo, gaz: Gazetteer) -> tuple[list[str], str | None]:
    """All LGUs the RDO covers, and its province if it is in scope."""
    raw: list[str] = []
    for d in fi.dos.values():
        for t in d.lgus_text:
            raw += split_lgus(t)
    if fi.composed_of == 'lgus' and not raw:
        raw = split_lgus(fi.composition)
    if raw:
        names, provs = resolve_lgus(raw, gaz)
        prov, n = provs.most_common(1)[0]
        in_scope = prov if prov in IN_SCOPE_PROVINCES and n / sum(provs.values()) > 0.5 else None
        return names, in_scope
    # barangay-composed RDOs (NCR): the city is named in the RDO name or the newest DO's subject
    subject = re.sub(r'REVENUE\s+REGION\s+(?:NO\.?\s*)?\d+[A-C]?\s*[-–,]?\s*[A-Za-z .]+?(?:CITY)?(?=[,.;]|\s+(?:FOR|AFTER|WITHIN)\b|$)',
                     ' ', getattr(fi, 'newest_subject', ''), flags=re.I)
    cities = gaz.ncr_cities_in(' '.join([fi.rdo_name, subject]))
    return cities, ('Metro Manila' if cities else None)


def do_lgus(d: DOInfo, fi_lgus: list[str], gaz: Gazetteer) -> list[str]:
    if not d.lgus_text:
        return fi_lgus
    raw = [x for t in d.lgus_text for x in split_lgus(t)]
    return resolve_lgus(raw, gaz)[0]


# ---------------------------------------------------------------- main

def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def main() -> None:
    gaz = Gazetteer()
    log: list[str] = [f'# inventory run {datetime.now().isoformat(timespec="seconds")}']
    known = {sha256(p): p for p in BIR.glob('RDO-*/*') if p.is_file()}

    # 1. file new inbox files
    inbox = sorted(p for p in INBOX.rglob('*') if p.is_file() and not p.name.startswith('.'))
    for p in inbox:
        rel = p.relative_to(INBOX)
        digest = sha256(p)
        if digest in known:
            DUPLICATES.mkdir(exist_ok=True)
            dest = DUPLICATES / p.name
            shutil.move(p, dest)
            log.append(f'DUPLICATE  {rel}  identical to {known[digest].relative_to(BIR)}; moved to {dest.relative_to(BIR)}')
            continue
        if p.suffix.lower() not in EXCEL_SUFFIXES:
            UNREADABLE.mkdir(exist_ok=True)
            shutil.move(p, UNREADABLE / p.name)
            log.append(f'UNREADABLE {rel}  not an Excel file; moved to _unreadable/')
            continue
        try:
            fi = read_file(p, gaz)
        except Exception as e:  # noqa: BLE001 - any failure means a human should look at the file
            UNREADABLE.mkdir(exist_ok=True)
            shutil.move(p, UNREADABLE / p.name)
            log.append(f'UNREADABLE {rel}  {type(e).__name__}: {e}; moved to _unreadable/')
            continue
        dest_dir = BIR / f'RDO-{fi.rdo}'
        dest_dir.mkdir(exist_ok=True)
        dest = dest_dir / p.name
        if dest.exists():
            log.append(f'CONFLICT   {rel}  {dest.relative_to(BIR)} exists with different contents; left in inbox')
            continue
        shutil.move(p, dest)
        known[digest] = dest
        log.append(f'MOVED      {rel}  -> {dest.relative_to(BIR)}')
    for d in sorted((x for x in INBOX.rglob('*') if x.is_dir()), reverse=True):
        if not any(d.iterdir()):
            d.rmdir()  # empty subfolders left behind by the move

    # 2. rebuild the inventory from everything filed
    rows, sheets = [], []
    for p in sorted(BIR.glob('RDO-*/*')):
        if not p.is_file() or p.name.startswith('.'):
            continue
        fi = read_file(p, gaz)
        if p.parent.name != f'RDO-{fi.rdo}':
            fi.notes.append(f'file sits in {p.parent.name} but its contents say RDO {fi.rdo}')
        lgus, prov = file_lgus(fi, gaz)
        for n in fi.notes:
            log.append(f'NOTE       RDO-{fi.rdo} {p.name}: {n}')
        for d in sorted(fi.dos.values(), key=lambda d: min(d.froms) if d.froms else date.max):
            rows.append({
                'rdo': f'RDO-{fi.rdo}', 'rdo_name': fi.rdo_name, 'cities_covered': '; '.join(do_lgus(d, lgus, gaz)),
                'department_order': d.do, 'effective_date': min(d.froms).isoformat() if d.froms else '',
                'file_name': p.name, 'file_type': p.suffix.lower().lstrip('.'), 'in_scope': 'yes' if prov else 'no',
            })
            for s in d.sheets:
                sr = next(x for x in fi.sheet_rows if x['sheet'] == s and x['do'] == d.do)
                sheets.append({'rdo': f'RDO-{fi.rdo}', 'file_name': p.name, 'department_order': d.do, 'sheet': s,
                               'revisions': '; '.join(d.revisions), 'lgus_in_notice': ' | '.join(d.lgus_text),
                               'do_from_header': sr['do_from_header'], 'do_from_tab': sr['do_from_tab'],
                               'rdo_in_subject': sr['rdo_in_subject'], 'subject': sr['subject']})

    rows.sort(key=lambda r: (r['rdo'], r['effective_date'] or '9999'))
    with open(BIR / 'inventory.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    REVIEW.mkdir(exist_ok=True)
    with open(REVIEW / 'bir_sheets.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(sheets[0]))
        w.writeheader()
        w.writerows(sheets)
    # moves, duplicates and unreadable files are history, appended; notes describe the current files, rewritten
    filed = [l for l in log[1:] if not l.startswith('NOTE')]
    if filed:
        with open(BIR / 'filing_log.txt', 'a', encoding='utf-8') as f:
            f.write('\n'.join([log[0]] + filed) + '\n\n')
    (BIR / 'inventory_log.txt').write_text('\n'.join(l for l in log if l.startswith('NOTE')) + '\n', encoding='utf-8')

    rdos = defaultdict(set)
    for r in rows:
        rdos[(r['rdo'], r['in_scope'])].add(r['department_order'])
    print(f"files: {len({r['file_name'] for r in rows})}  RDOs: {len({r['rdo'] for r in rows})}  "
          f"in scope: {len({k for k, s in rdos if s == 'yes'})}  DO rows: {len(rows)}")
    print('\n'.join(l for l in log if not l.startswith('NOTE')))


if __name__ == '__main__':
    main()
